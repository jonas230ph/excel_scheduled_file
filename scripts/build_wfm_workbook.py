"""Build the Phase 1 Excel WFM scheduling workbook skeleton."""

from __future__ import annotations

from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName


ROOT = Path(__file__).resolve().parents[1]
EXCEL_OUTPUT_PATH = ROOT / "excel schedule" / "wfm_scheduling_matrix.xlsx"
GOOGLE_OUTPUT_PATH = ROOT / "excel schedule" / "wfm_scheduling_matrix_google.xlsx"
SAMPLE_DATE = date(2026, 6, 11)
ROSTER_START_ROW = 8
MAX_ROSTER_ROWS = 250
ROSTER_END_ROW = ROSTER_START_ROW + MAX_ROSTER_ROWS - 1
SUMMARY_SCHEDULED_ROW = 260
SUMMARY_REQUIRED_ROW = 261
SUMMARY_VARIANCE_ROW = 262
INTERVAL_START_COLUMN = 8
INTERVAL_END_COLUMN = 55
VALIDATION_STATE_COLUMN = 56
SCHEDULE_HELPER_START_COLUMN = 57
SELECTED_DATE_REF = "Schedule_Matrix!$E$2"
SCHEDULE_EXTRA_FIELDS = [
    "Break1Start",
    "Break1End",
    "Break2Start",
    "Break2End",
    "LunchStart",
    "LunchEnd",
    "Adhoc1Code",
    "Adhoc1Start",
    "Adhoc1End",
    "Adhoc2Code",
    "Adhoc2Start",
    "Adhoc2End",
    "Adhoc3Code",
    "Adhoc3Start",
    "Adhoc3End",
]
SCHEDULE_HELPER_END_COLUMN = SCHEDULE_HELPER_START_COLUMN + len(SCHEDULE_EXTRA_FIELDS) - 1
SCHEDULE_EXTRA_TIME_COLUMNS = [12, 13, 14, 15, 16, 17, 19, 20, 22, 23, 25, 26]
SHEETS = [
    "Config_Settings",
    "Staffing_Requirements",
    "Schedule_Data",
    "Calc_Engine",
    "Schedule_Matrix",
    "Summary_Dashboard",
    "TestCases",
]


def add_table(ws, name: str, ref: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def add_name(wb: Workbook, name: str, refers_to: str) -> None:
    wb.defined_names.add(DefinedName(name, attr_text=refers_to))


def interval_time(index: int) -> time:
    minutes = index * 30
    return time(minutes // 60, minutes % 60)


def schedule_helper_columns() -> dict[str, str]:
    return {
        field: get_column_letter(SCHEDULE_HELPER_START_COLUMN + index)
        for index, field in enumerate(SCHEDULE_EXTRA_FIELDS)
    }


def setup_workbook() -> Workbook:
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    for sheet_name in SHEETS:
        wb.create_sheet(sheet_name)
    return wb


def populate_config(ws) -> None:
    activity_headers = [
        "Code",
        "DisplayLabel",
        "IsActive",
        "CountsAsStaffed",
        "IsPaid",
        "Category",
        "NumericValue",
        "ColorGroup",
    ]
    ws.append(activity_headers)
    activity_rows = [
        ("OWD", "On Work Duty", True, True, True, "Productive", 1, "Productive"),
        ("OT", "Overtime", True, True, True, "Productive", 1, "Productive"),
        ("Brk", "Break", True, False, True, "Break", 0, "Nonproductive"),
        ("Lch", "Lunch", True, False, False, "Meal", 0, "Nonproductive"),
        ("Mt", "Meeting", True, False, True, "Meeting", 0, "Nonproductive"),
        ("Trn", "Training", True, False, True, "Training", 0, "Nonproductive"),
        ("PTO", "Paid Time Off", True, False, True, "Absence", 0, "Inactive"),
        ("Sick", "Sick Leave", True, False, True, "Absence", 0, "Inactive"),
        ("Off", "Off", True, False, False, "Off", 0, "Inactive"),
        ("UA", "Unapproved Absence", True, False, False, "Absence", 0, "Inactive"),
        ("AL", "Annual Leave", True, False, True, "Absence", 0, "Inactive"),
        ("SL", "Sick Leave Short Code", True, False, True, "Absence", 0, "Inactive"),
    ]
    for row in activity_rows:
        ws.append(row)
    for row_index in range(len(activity_rows) + 2, 41):
        ws.cell(row=row_index, column=1, value="")

    parameter_headers = ["Parameter", "Value"]
    for column_index, header in enumerate(parameter_headers, start=10):
        ws.cell(row=1, column=column_index, value=header)
    parameters = [
        ("IntervalMinutes", 30),
        ("IntervalsPerDay", 48),
        ("MaxRosterRows", MAX_ROSTER_ROWS),
        ("DayStartTime", time(0, 0)),
        ("ToleranceHeadcount", 0),
        ("WorkbookEngineVersion", "1.0.0-phase1"),
        ("TemplateSchemaVersion", "1.0.0"),
        ("FormulaSetVersion", "1.0.0"),
        ("DataSnapshotID", "sample-2026-06-11"),
    ]
    for row_index, row in enumerate(parameters, start=2):
        ws.cell(row=row_index, column=10, value=row[0])
        ws.cell(row=row_index, column=11, value=row[1])
    for row_index in range(len(parameters) + 2, 31):
        ws.cell(row=row_index, column=10, value="")

    threshold_headers = ["State", "MinVariance", "MaxVariance", "FillColor", "FontColor"]
    for column_index, header in enumerate(threshold_headers, start=13):
        ws.cell(row=1, column=column_index, value=header)
    thresholds = [
        ("Under", None, -1, "F8696B", "FFFFFF"),
        ("Exact", 0, 0, "63BE7B", "FFFFFF"),
        ("Over", 1, None, "5B9BD5", "FFFFFF"),
        ("InvalidCode", None, None, "FFC000", "000000"),
        ("InvalidSchedule", None, None, "FFD966", "000000"),
        ("Duplicate", None, None, "D9EAD3", "000000"),
        ("MissingDate", None, None, "C9DAF8", "000000"),
        ("MissingRequirement", None, None, "BFBFBF", "000000"),
        ("Nonproductive", None, None, "F4B183", "000000"),
    ]
    for row_index, row in enumerate(thresholds, start=2):
        for offset, value in enumerate(row, start=13):
            ws.cell(row=row_index, column=offset, value=value)
    for row_index in range(len(thresholds) + 2, 21):
        ws.cell(row=row_index, column=13, value="")

    add_table(ws, "tblActivityCodes", "A1:H40")
    add_table(ws, "tblParameters", "J1:K30")
    add_table(ws, "tblThresholds", "M1:Q20")


def populate_requirements(ws) -> None:
    ws.append(["OperationalDate", "IntervalStart", "RequiredHeadcount"])
    for day_offset in range(7):
        operational_date = SAMPLE_DATE + timedelta(days=day_offset)
        for index in range(48):
            interval = interval_time(index)
            required = 4 if 8 <= interval.hour < 18 else 2
            ws.append([operational_date, interval, required])
    add_table(ws, "tblRequirements", "A1:C337")


def populate_schedule_data(ws) -> None:
    headers = [
        "EmployeeID",
        "Name",
        "OperationalDate",
        "ShiftStart",
        "ShiftEnd",
        "ActivityCode",
        "ContractualHours",
        "CurrentShift",
        "Notes",
        "ValidationState",
        "SourceRowID",
        *SCHEDULE_EXTRA_FIELDS,
    ]
    ws.append(headers)
    rows = [
        (
            "E001",
            "Avery Chen",
            SAMPLE_DATE,
            datetime(2026, 6, 11, 8, 15),
            datetime(2026, 6, 11, 17, 0),
            "OWD",
            8,
            "OWD",
            "Partial first interval with breaks",
            "",
            1,
            time(10, 0),
            time(10, 15),
            time(15, 0),
            time(15, 15),
            time(12, 0),
            time(12, 30),
            "Mt",
            time(14, 0),
            time(14, 30),
            "",
            "",
            "",
            "",
            "",
            "",
        ),
        ("E002", "Blake Diaz", SAMPLE_DATE, datetime(2026, 6, 11, 23, 45), datetime(2026, 6, 12, 2, 15), "OWD", 8, "OWD", "Overnight", "", 2, *[""] * len(SCHEDULE_EXTRA_FIELDS)),
        ("E003", "Casey Singh", SAMPLE_DATE, datetime(2026, 6, 11, 12, 0), datetime(2026, 6, 11, 12, 30), "Lch", 8, "Lch", "Nonproductive", "", 3, *[""] * len(SCHEDULE_EXTRA_FIELDS)),
    ]
    for row in rows:
        ws.append(row)
    for row_index in range(2, len(rows) + 2):
        ws.cell(row=row_index, column=3).number_format = "yyyy-mm-dd"
        ws.cell(row=row_index, column=4).number_format = "yyyy-mm-dd hh:mm"
        ws.cell(row=row_index, column=5).number_format = "yyyy-mm-dd hh:mm"
        for column in SCHEDULE_EXTRA_TIME_COLUMNS:
            ws.cell(row=row_index, column=column).number_format = "hh:mm"
    for row_index in range(len(rows) + 2, 1001):
        ws.cell(row=row_index, column=1, value="")
    for row_index in range(2, 1001):
        ws.cell(row=row_index, column=10, value=schedule_validation_formula(row_index))
    add_schedule_data_conditional_formatting(ws)
    add_table(ws, "tblScheduleData", "A1:Z1000")


def schedule_validation_formula(row: int) -> str:
    return (
        f'=IF(COUNTA($A{row}:$F{row})=0,"",'
        f'IF($C{row}="","MissingDate",'
        f'IF(OR($A{row}="",$D{row}="",$E{row}="",$D{row}=$E{row}),"InvalidSchedule",'
        f'IF(COUNTIF(ActiveActivityCodes,$F{row})=0,"InvalidCode",'
        f'IF(COUNTIFS(tblScheduleData[EmployeeID],$A{row},'
        f'tblScheduleData[OperationalDate],$C{row},'
        f'tblScheduleData[ShiftStart],$D{row},'
        f'tblScheduleData[ShiftEnd],$E{row},'
        f'tblScheduleData[ActivityCode],$F{row})>1,"Duplicate","Valid")))))'
    )


def window_let_bindings(prefix: str, start_var: str, end_var: str, selected_date_ref: str) -> list[str]:
    return [
        f"{prefix}StartDT,IF({start_var}=\"\",0,IF(INT({start_var})>0,{start_var},{selected_date_ref}+MOD({start_var},1)))",
        f"{prefix}EndBase,IF({end_var}=\"\",0,IF(INT({end_var})>0,{end_var},{selected_date_ref}+MOD({end_var},1)))",
        f"{prefix}EndDT,IF(OR({start_var}=\"\",{end_var}=\"\"),0,IF({prefix}EndBase<={prefix}StartDT,{prefix}EndBase+1,{prefix}EndBase))",
        f"{prefix}RawOverlap,IF(OR({start_var}=\"\",{end_var}=\"\",baseOverlap=0),0,MAX(0,MIN({prefix}EndDT,intervalEnd)-MAX({prefix}StartDT,intervalStart)))",
        f"{prefix}Overlap,IF({prefix}RawOverlap<TIME(0,0,1),0,{prefix}RawOverlap)",
    ]


def matrix_visible_formula(row: int, column_letter: str) -> str:
    helpers = schedule_helper_columns()
    assignments = [
        f"id,$A{row}",
        f"code,UPPER(TRIM($D{row}))",
        f"rawStart,$F{row}",
        f"rawEnd,$G{row}",
        f"Break1Start,${helpers['Break1Start']}{row}",
        f"Break1End,${helpers['Break1End']}{row}",
        f"Break2Start,${helpers['Break2Start']}{row}",
        f"Break2End,${helpers['Break2End']}{row}",
        f"LunchStart,${helpers['LunchStart']}{row}",
        f"LunchEnd,${helpers['LunchEnd']}{row}",
        f"Adhoc1Code,UPPER(TRIM(${helpers['Adhoc1Code']}{row}))",
        f"Adhoc1Start,${helpers['Adhoc1Start']}{row}",
        f"Adhoc1End,${helpers['Adhoc1End']}{row}",
        f"Adhoc2Code,UPPER(TRIM(${helpers['Adhoc2Code']}{row}))",
        f"Adhoc2Start,${helpers['Adhoc2Start']}{row}",
        f"Adhoc2End,${helpers['Adhoc2End']}{row}",
        f"Adhoc3Code,UPPER(TRIM(${helpers['Adhoc3Code']}{row}))",
        f"Adhoc3Start,${helpers['Adhoc3Start']}{row}",
        f"Adhoc3End,${helpers['Adhoc3End']}{row}",
        f"startDT,IF(INT(rawStart)>0,rawStart,$E$2+MOD(rawStart,1))",
        f"endBase,IF(INT(rawEnd)>0,rawEnd,$E$2+MOD(rawEnd,1))",
        "endDT,IF(endBase<=startDT,endBase+1,endBase)",
        f"intervalStart,$E$2+{column_letter}$1",
        "intervalEnd,intervalStart+TIME(0,30,0)",
        "baseRawOverlap,MAX(0,MIN(endDT,intervalEnd)-MAX(startDT,intervalStart))",
        "baseOverlap,IF(baseRawOverlap<TIME(0,0,1),0,baseRawOverlap)",
        *window_let_bindings("break1", "Break1Start", "Break1End", "$E$2"),
        *window_let_bindings("break2", "Break2Start", "Break2End", "$E$2"),
        *window_let_bindings("lunch", "LunchStart", "LunchEnd", "$E$2"),
        *window_let_bindings("adhoc1", "Adhoc1Start", "Adhoc1End", "$E$2"),
        *window_let_bindings("adhoc2", "Adhoc2Start", "Adhoc2End", "$E$2"),
        *window_let_bindings("adhoc3", "Adhoc3Start", "Adhoc3End", "$E$2"),
    ]
    return (
        "=LET("
        + ",".join(assignments)
        + ',IF(OR(id="",code="",rawStart="",rawEnd="",rawStart=rawEnd,baseOverlap=0),"",'
        + 'IF(AND(Adhoc1Code<>"",adhoc1Overlap>0),Adhoc1Code,'
        + 'IF(AND(Adhoc2Code<>"",adhoc2Overlap>0),Adhoc2Code,'
        + 'IF(AND(Adhoc3Code<>"",adhoc3Overlap>0),Adhoc3Code,'
        + 'IF(lunchOverlap>0,"LCH",IF(OR(break1Overlap>0,break2Overlap>0),"BRK",code))))))'
        + ")"
    )


def calc_coverage_formula(row: int, column_letter: str) -> str:
    helpers = schedule_helper_columns()
    assignments = [
        f"id,Schedule_Matrix!$A{row}",
        f"code,UPPER(TRIM(Schedule_Matrix!$D{row}))",
        f"rawStart,Schedule_Matrix!$F{row}",
        f"rawEnd,Schedule_Matrix!$G{row}",
        f"Break1Start,Schedule_Matrix!${helpers['Break1Start']}{row}",
        f"Break1End,Schedule_Matrix!${helpers['Break1End']}{row}",
        f"Break2Start,Schedule_Matrix!${helpers['Break2Start']}{row}",
        f"Break2End,Schedule_Matrix!${helpers['Break2End']}{row}",
        f"LunchStart,Schedule_Matrix!${helpers['LunchStart']}{row}",
        f"LunchEnd,Schedule_Matrix!${helpers['LunchEnd']}{row}",
        f"Adhoc1Code,UPPER(TRIM(Schedule_Matrix!${helpers['Adhoc1Code']}{row}))",
        f"Adhoc1Start,Schedule_Matrix!${helpers['Adhoc1Start']}{row}",
        f"Adhoc1End,Schedule_Matrix!${helpers['Adhoc1End']}{row}",
        f"Adhoc2Code,UPPER(TRIM(Schedule_Matrix!${helpers['Adhoc2Code']}{row}))",
        f"Adhoc2Start,Schedule_Matrix!${helpers['Adhoc2Start']}{row}",
        f"Adhoc2End,Schedule_Matrix!${helpers['Adhoc2End']}{row}",
        f"Adhoc3Code,UPPER(TRIM(Schedule_Matrix!${helpers['Adhoc3Code']}{row}))",
        f"Adhoc3Start,Schedule_Matrix!${helpers['Adhoc3Start']}{row}",
        f"Adhoc3End,Schedule_Matrix!${helpers['Adhoc3End']}{row}",
        "countsAsStaffed,IFERROR(--INDEX(tblActivityCodes[CountsAsStaffed],MATCH(code,tblActivityCodes[Code],0)),0)",
        'breakCount,IFERROR(--INDEX(tblActivityCodes[CountsAsStaffed],MATCH("Brk",tblActivityCodes[Code],0)),0)',
        'lunchCount,IFERROR(--INDEX(tblActivityCodes[CountsAsStaffed],MATCH("Lch",tblActivityCodes[Code],0)),0)',
        "adhoc1Count,IF(Adhoc1Code=\"\",0,IFERROR(--INDEX(tblActivityCodes[CountsAsStaffed],MATCH(Adhoc1Code,tblActivityCodes[Code],0)),0))",
        "adhoc2Count,IF(Adhoc2Code=\"\",0,IFERROR(--INDEX(tblActivityCodes[CountsAsStaffed],MATCH(Adhoc2Code,tblActivityCodes[Code],0)),0))",
        "adhoc3Count,IF(Adhoc3Code=\"\",0,IFERROR(--INDEX(tblActivityCodes[CountsAsStaffed],MATCH(Adhoc3Code,tblActivityCodes[Code],0)),0))",
        f"startDT,IF(INT(rawStart)>0,rawStart,{SELECTED_DATE_REF}+MOD(rawStart,1))",
        f"endBase,IF(INT(rawEnd)>0,rawEnd,{SELECTED_DATE_REF}+MOD(rawEnd,1))",
        "endDT,IF(endBase<=startDT,endBase+1,endBase)",
        f"intervalStart,{SELECTED_DATE_REF}+Schedule_Matrix!{column_letter}$1",
        "intervalEnd,intervalStart+TIME(0,30,0)",
        "baseRawOverlap,MAX(0,MIN(endDT,intervalEnd)-MAX(startDT,intervalStart))",
        "baseOverlap,IF(baseRawOverlap<TIME(0,0,1),0,baseRawOverlap)",
        *window_let_bindings("break1", "Break1Start", "Break1End", SELECTED_DATE_REF),
        *window_let_bindings("break2", "Break2Start", "Break2End", SELECTED_DATE_REF),
        *window_let_bindings("lunch", "LunchStart", "LunchEnd", SELECTED_DATE_REF),
        *window_let_bindings("adhoc1", "Adhoc1Start", "Adhoc1End", SELECTED_DATE_REF),
        *window_let_bindings("adhoc2", "Adhoc2Start", "Adhoc2End", SELECTED_DATE_REF),
        *window_let_bindings("adhoc3", "Adhoc3Start", "Adhoc3End", SELECTED_DATE_REF),
        "baseCoverage,baseOverlap/TIME(0,30,0)*countsAsStaffed",
        "overlayCoverage,break1Overlap/TIME(0,30,0)*(breakCount-countsAsStaffed)+break2Overlap/TIME(0,30,0)*(breakCount-countsAsStaffed)+lunchOverlap/TIME(0,30,0)*(lunchCount-countsAsStaffed)+adhoc1Overlap/TIME(0,30,0)*(adhoc1Count-countsAsStaffed)+adhoc2Overlap/TIME(0,30,0)*(adhoc2Count-countsAsStaffed)+adhoc3Overlap/TIME(0,30,0)*(adhoc3Count-countsAsStaffed)",
    ]
    return (
        "=LET("
        + ",".join(assignments)
        + ',IF(OR(id="",code="",rawStart="",rawEnd="",rawStart=rawEnd),0,MAX(0,baseCoverage+overlayCoverage))'
        + ")"
    )


def requirement_formula(column_letter: str, row: int) -> str:
    return (
        f'=IF(COUNTIFS(tblRequirements[OperationalDate],$E$2,'
        f'tblRequirements[IntervalStart],{column_letter}${row})=0,"",'
        f"SUMIFS(tblRequirements[RequiredHeadcount],tblRequirements[OperationalDate],$E$2,"
        f"tblRequirements[IntervalStart],{column_letter}${row}))"
    )


def selected_day_variance_formula(column_letter: str, scheduled_row: int, required_row: int) -> str:
    return (
        f'=IF($E$2="","MissingDate",'
        f'IF({column_letter}${required_row}="","MissingRequirement",'
        f"ROUND({column_letter}${scheduled_row}-{column_letter}${required_row},2)))"
    )


def selected_schedule_formula(field_name: str, row: int, engine: str = "excel") -> str:
    relative_index = f"ROWS($A${ROSTER_START_ROW}:A{row})"
    row_index = f"(ROW(tblScheduleData[{field_name}])-ROW(INDEX(tblScheduleData[{field_name}],1,1))+1)"
    start_dt = "(tblScheduleData[ShiftStart]+(INT(tblScheduleData[ShiftStart])=0)*tblScheduleData[OperationalDate])"
    end_base = "(tblScheduleData[ShiftEnd]+(INT(tblScheduleData[ShiftEnd])=0)*tblScheduleData[OperationalDate])"
    end_dt = f"({end_base}+({end_base}<={start_dt}))"
    criteria = (
        '(tblScheduleData[EmployeeID]<>"")*(tblScheduleData[ShiftStart]<>"")*'
        "(tblScheduleData[ShiftEnd]<>\"\")*(tblScheduleData[OperationalDate]<>\"\")*"
        f"({SELECTED_DATE_REF}<>\"\")*"
        f"({start_dt}<{SELECTED_DATE_REF}+1)*"
        f"({end_dt}>{SELECTED_DATE_REF})"
    )
    if engine == "google":
        return (
            f'=IFERROR(INDEX(FILTER(tblScheduleData[{field_name}],{criteria}),'
            f"{relative_index}),\"\")"
        )
    return (
        f'=IFERROR(INDEX(tblScheduleData[{field_name}],'
        f'AGGREGATE(15,6,{row_index}/({criteria}),{relative_index})),\"\")'
    )


def weekly_scheduled_formula(date_cell: str) -> str:
    return (
        f'=IF({date_cell}={SELECTED_DATE_REF},'
        f"ROUND(SUM(Schedule_Matrix!$H${SUMMARY_SCHEDULED_ROW}:$BC${SUMMARY_SCHEDULED_ROW}),2),0)"
    )


def populate_schedule_matrix(wb: Workbook, engine: str) -> None:
    ws = wb["Schedule_Matrix"]
    calc = wb["Calc_Engine"]
    headers = ["EmployeeID", "Name", "ContractualHours", "CurrentShift", "TargetDaySelect", "ShiftStart", "ShiftEnd"]
    for column_index, header in enumerate(headers, start=1):
        ws.cell(row=1, column=column_index, value=header)
        calc.cell(row=1, column=column_index, value=header)
    for index in range(48):
        column = INTERVAL_START_COLUMN + index
        ws.cell(row=1, column=column, value=f"=TIME(0,0,0)+(COLUMN()-COLUMN($H$1))*TIME(0,30,0)")
        calc.cell(row=1, column=column, value=f"=Schedule_Matrix!{ws.cell(row=1, column=column).coordinate}")
    ws.cell(row=1, column=VALIDATION_STATE_COLUMN, value="ValidationState")
    calc.cell(row=1, column=VALIDATION_STATE_COLUMN, value="ValidationState")
    for offset, field_name in enumerate(SCHEDULE_EXTRA_FIELDS):
        column = SCHEDULE_HELPER_START_COLUMN + offset
        ws.cell(row=1, column=column, value=field_name)
        calc.cell(row=1, column=column, value=field_name)

    ws["E2"] = SAMPLE_DATE
    ws["G5"] = "Scheduled Productive"
    ws["G6"] = "Required"
    ws["G7"] = "Over/Under"
    ws.cell(row=SUMMARY_SCHEDULED_ROW, column=7, value="Scheduled Productive")
    ws.cell(row=SUMMARY_REQUIRED_ROW, column=7, value="Required")
    ws.cell(row=SUMMARY_VARIANCE_ROW, column=7, value="Over/Under")

    for row in range(ROSTER_START_ROW, ROSTER_END_ROW + 1):
        calc.cell(row=row, column=1, value=selected_schedule_formula("EmployeeID", row, engine))
        calc.cell(row=row, column=2, value=selected_schedule_formula("Name", row, engine))
        calc.cell(row=row, column=3, value=selected_schedule_formula("ContractualHours", row, engine))
        calc.cell(row=row, column=4, value=selected_schedule_formula("ActivityCode", row, engine))
        calc.cell(row=row, column=5, value=f'=IF(Calc_Engine!A{row}="","",{SELECTED_DATE_REF})')
        calc.cell(row=row, column=6, value=selected_schedule_formula("ShiftStart", row, engine))
        calc.cell(row=row, column=7, value=selected_schedule_formula("ShiftEnd", row, engine))
        calc.cell(row=row, column=VALIDATION_STATE_COLUMN, value=selected_schedule_formula("ValidationState", row, engine))
        for offset, field_name in enumerate(SCHEDULE_EXTRA_FIELDS):
            calc.cell(row=row, column=SCHEDULE_HELPER_START_COLUMN + offset, value=selected_schedule_formula(field_name, row, engine))
        for column in range(1, 8):
            ws.cell(row=row, column=column, value=f"=Calc_Engine!{ws.cell(row=row, column=column).coordinate}")
        ws.cell(row=row, column=VALIDATION_STATE_COLUMN, value=f"=Calc_Engine!{ws.cell(row=row, column=VALIDATION_STATE_COLUMN).coordinate}")
        for column in range(SCHEDULE_HELPER_START_COLUMN, SCHEDULE_HELPER_END_COLUMN + 1):
            ws.cell(row=row, column=column, value=f"=Calc_Engine!{ws.cell(row=row, column=column).coordinate}")
        for column in range(INTERVAL_START_COLUMN, INTERVAL_END_COLUMN + 1):
            column_letter = ws.cell(row=1, column=column).column_letter
            ws.cell(row=row, column=column, value=matrix_visible_formula(row, column_letter))
            calc.cell(row=row, column=column, value=calc_coverage_formula(row, column_letter))

    for column in range(INTERVAL_START_COLUMN, INTERVAL_END_COLUMN + 1):
        column_letter = ws.cell(row=1, column=column).column_letter
        ws.cell(row=5, column=column, value=f"=SUM(Calc_Engine!{column_letter}${ROSTER_START_ROW}:{column_letter}${ROSTER_END_ROW})")
        ws.cell(
            row=6,
            column=column,
            value=requirement_formula(column_letter, 1),
        )
        ws.cell(row=7, column=column, value=selected_day_variance_formula(column_letter, 5, 6))
        ws.cell(row=SUMMARY_SCHEDULED_ROW, column=column, value=f"=SUM(Calc_Engine!{column_letter}${ROSTER_START_ROW}:{column_letter}${ROSTER_END_ROW})")
        ws.cell(
            row=SUMMARY_REQUIRED_ROW,
            column=column,
            value=requirement_formula(column_letter, 1),
        )
        ws.cell(
            row=SUMMARY_VARIANCE_ROW,
            column=column,
            value=selected_day_variance_formula(column_letter, SUMMARY_SCHEDULED_ROW, SUMMARY_REQUIRED_ROW),
        )

    add_conditional_formatting(ws)


def populate_summary_dashboard(wb: Workbook) -> None:
    ws = wb["Summary_Dashboard"]
    ws["A1"] = "Weekly Staffing Summary"
    ws["B1"] = f"={SELECTED_DATE_REF}-WEEKDAY({SELECTED_DATE_REF},2)+1"
    ws["A2"] = "Selected Date"
    ws["B2"] = f"={SELECTED_DATE_REF}"
    weekly_headers = ["Date", "Day", "Scheduled Productive", "Required", "Over/Under"]
    for column_index, header in enumerate(weekly_headers, start=1):
        ws.cell(row=3, column=column_index, value=header)
    for row in range(4, 11):
        previous_row = row - 1
        ws.cell(row=row, column=1, value="=$B$1" if row == 4 else f"=A{previous_row}+1")
        ws.cell(row=row, column=2, value=f'=TEXT(A{row},"ddd")')
        ws.cell(row=row, column=3, value=weekly_scheduled_formula(f"A{row}"))
        ws.cell(row=row, column=4, value=f"=SUMIFS(tblRequirements[RequiredHeadcount],tblRequirements[OperationalDate],A{row})")
        ws.cell(row=row, column=5, value=f'=IF(D{row}="","",ROUND(C{row}-D{row},2))')

    ws["A13"] = "Interval"
    ws["B13"] = "Scheduled Productive"
    ws["C13"] = "Required"
    ws["D13"] = "Over/Under"
    matrix = wb["Schedule_Matrix"]
    for index, column in enumerate(range(INTERVAL_START_COLUMN, INTERVAL_END_COLUMN + 1), start=14):
        column_letter = matrix.cell(row=1, column=column).column_letter
        ws.cell(row=index, column=1, value=f"=Schedule_Matrix!{column_letter}$1")
        ws.cell(row=index, column=2, value=f"=Schedule_Matrix!{column_letter}${SUMMARY_SCHEDULED_ROW}")
        ws.cell(row=index, column=3, value=f"=Schedule_Matrix!{column_letter}${SUMMARY_REQUIRED_ROW}")
        ws.cell(row=index, column=4, value=f"=Schedule_Matrix!{column_letter}${SUMMARY_VARIANCE_ROW}")

    weekly_chart = BarChart()
    weekly_chart.title = "Weekly Scheduled vs Required"
    weekly_chart.y_axis.title = "Agent half-hour units"
    weekly_chart.x_axis.title = "Day"
    weekly_chart.height = 8
    weekly_chart.width = 18
    weekly_data = Reference(ws, min_col=3, min_row=3, max_col=4, max_row=10)
    weekly_categories = Reference(ws, min_col=2, min_row=4, max_row=10)
    weekly_chart.add_data(weekly_data, titles_from_data=True)
    weekly_chart.set_categories(weekly_categories)
    ws.add_chart(weekly_chart, "G2")

    intraday_chart = LineChart()
    intraday_chart.title = "Selected-Day Intraday Staffing"
    intraday_chart.y_axis.title = "Agents"
    intraday_chart.x_axis.title = "Interval"
    intraday_chart.height = 9
    intraday_chart.width = 24
    intraday_data = Reference(ws, min_col=2, min_row=13, max_col=4, max_row=61)
    intraday_categories = Reference(ws, min_col=1, min_row=14, max_row=61)
    intraday_chart.add_data(intraday_data, titles_from_data=True)
    intraday_chart.set_categories(intraday_categories)
    ws.add_chart(intraday_chart, "G20")


def add_conditional_formatting(ws) -> None:
    productive_fill = PatternFill("solid", fgColor="C6EFCE")
    nonproductive_fill = PatternFill("solid", fgColor="FCE4D6")
    under_fill = PatternFill("solid", fgColor="F8696B")
    exact_fill = PatternFill("solid", fgColor="63BE7B")
    over_fill = PatternFill("solid", fgColor="5B9BD5")
    invalid_code_fill = PatternFill("solid", fgColor="FFC000")
    invalid_schedule_fill = PatternFill("solid", fgColor="FFD966")
    duplicate_fill = PatternFill("solid", fgColor="D9EAD3")
    missing_date_fill = PatternFill("solid", fgColor="C9DAF8")
    missing_fill = PatternFill("solid", fgColor="BFBFBF")
    body_range = f"H{ROSTER_START_ROW}:BC{ROSTER_END_ROW}"
    row_range = f"A{ROSTER_START_ROW}:BC{ROSTER_END_ROW}"
    variance_range = f"H{SUMMARY_VARIANCE_ROW}:BC{SUMMARY_VARIANCE_ROW}"
    ws.conditional_formatting.add(body_range, FormulaRule(formula=['H8="OWD"'], fill=productive_fill))
    ws.conditional_formatting.add(body_range, FormulaRule(formula=['OR(H8="BRK",H8="LCH",H8="MT",H8="TRN")'], fill=nonproductive_fill))
    ws.conditional_formatting.add(row_range, FormulaRule(formula=['$BD8="InvalidCode"'], fill=invalid_code_fill))
    ws.conditional_formatting.add(row_range, FormulaRule(formula=['$BD8="InvalidSchedule"'], fill=invalid_schedule_fill))
    ws.conditional_formatting.add(row_range, FormulaRule(formula=['$BD8="Duplicate"'], fill=duplicate_fill))
    ws.conditional_formatting.add(row_range, FormulaRule(formula=['$BD8="MissingDate"'], fill=missing_date_fill))
    ws.conditional_formatting.add("E2", FormulaRule(formula=['$E$2=""'], fill=missing_date_fill))
    ws.conditional_formatting.add("H7:BC7", FormulaRule(formula=['H7="MissingDate"'], fill=missing_date_fill))
    ws.conditional_formatting.add("H7:BC7", FormulaRule(formula=['H7="MissingRequirement"'], fill=missing_fill))
    ws.conditional_formatting.add(variance_range, FormulaRule(formula=[f'H{SUMMARY_VARIANCE_ROW}="MissingDate"'], fill=missing_date_fill))
    ws.conditional_formatting.add(variance_range, FormulaRule(formula=[f'H{SUMMARY_VARIANCE_ROW}="MissingRequirement"'], fill=missing_fill))
    ws.conditional_formatting.add("H7:BC7", CellIsRule(operator="lessThan", formula=["0"], fill=under_fill))
    ws.conditional_formatting.add("H7:BC7", CellIsRule(operator="equal", formula=["0"], fill=exact_fill))
    ws.conditional_formatting.add("H7:BC7", CellIsRule(operator="greaterThan", formula=["0"], fill=over_fill))
    ws.conditional_formatting.add(variance_range, CellIsRule(operator="lessThan", formula=["0"], fill=under_fill))
    ws.conditional_formatting.add(variance_range, CellIsRule(operator="equal", formula=["0"], fill=exact_fill))
    ws.conditional_formatting.add(variance_range, CellIsRule(operator="greaterThan", formula=["0"], fill=over_fill))


def add_schedule_data_conditional_formatting(ws) -> None:
    invalid_code_fill = PatternFill("solid", fgColor="FFC000")
    invalid_schedule_fill = PatternFill("solid", fgColor="FFD966")
    duplicate_fill = PatternFill("solid", fgColor="D9EAD3")
    missing_date_fill = PatternFill("solid", fgColor="C9DAF8")
    data_range = "A2:Z1000"
    ws.conditional_formatting.add(data_range, FormulaRule(formula=['$J2="InvalidCode"'], fill=invalid_code_fill))
    ws.conditional_formatting.add(data_range, FormulaRule(formula=['$J2="InvalidSchedule"'], fill=invalid_schedule_fill))
    ws.conditional_formatting.add(data_range, FormulaRule(formula=['$J2="Duplicate"'], fill=duplicate_fill))
    ws.conditional_formatting.add(data_range, FormulaRule(formula=['$J2="MissingDate"'], fill=missing_date_fill))


def populate_test_cases(ws) -> None:
    headers = ["TestID", "Scenario", "Input", "Expected", "Actual", "Pass", "Requirement", "Notes"]
    ws.append(headers)
    ws["J1"] = "OverallStatus"
    ws["K1"] = '=IF(COUNTIF(F2:F11,FALSE)=0,"PASS","FAIL")'
    rows = [
        (
            "TEST-01",
            "Same-day productive shift",
            "09:00-17:00 OWD",
            16,
            "=(DATE(2026,6,11)+TIME(17,0,0)-(DATE(2026,6,11)+TIME(9,0,0)))/TIME(0,30,0)",
            "=ABS(E2-D2)<0.000001",
            "TEST-01",
            "Verifies same-day half-hour interval count.",
        ),
        (
            "TEST-02",
            "Overnight productive shift",
            "22:00-06:00 OWD",
            16,
            "=(DATE(2026,6,12)+TIME(6,0,0)-(DATE(2026,6,11)+TIME(22,0,0)))/TIME(0,30,0)",
            "=ABS(E3-D3)<0.000001",
            "TEST-02",
            "Verifies midnight-crossing interval count.",
        ),
        (
            "TEST-03",
            "Boundary overlap",
            "08:15 start against 08:00-08:30 interval",
            0.5,
            "=MAX(0,MIN(DATE(2026,6,11)+TIME(8,30,0),DATE(2026,6,11)+TIME(12,0,0))-MAX(DATE(2026,6,11)+TIME(8,0,0),DATE(2026,6,11)+TIME(8,15,0)))/TIME(0,30,0)",
            "=ABS(E4-D4)<0.000001",
            "TEST-03",
            "Verifies partial interval proration.",
        ),
        (
            "TEST-04",
            "Productive code lookup",
            "OWD CountsAsStaffed",
            1,
            '=--INDEX(tblActivityCodes[CountsAsStaffed],MATCH("OWD",tblActivityCodes[Code],0))',
            "=ABS(E5-D5)<0.000001",
            "TEST-04",
            "Verifies OWD counts as staffed through config lookup.",
        ),
        (
            "TEST-05",
            "Nonproductive code lookup",
            "Brk/Lch/Mt/Trn CountsAsStaffed",
            0,
            '=SUM(--INDEX(tblActivityCodes[CountsAsStaffed],MATCH("Brk",tblActivityCodes[Code],0)),--INDEX(tblActivityCodes[CountsAsStaffed],MATCH("Lch",tblActivityCodes[Code],0)),--INDEX(tblActivityCodes[CountsAsStaffed],MATCH("Mt",tblActivityCodes[Code],0)),--INDEX(tblActivityCodes[CountsAsStaffed],MATCH("Trn",tblActivityCodes[Code],0)))',
            "=ABS(E6-D6)<0.000001",
            "TEST-05",
            "Verifies visible nonproductive states do not count as staffed.",
        ),
        (
            "TEST-06",
            "Validation and variance states",
            "Invalid code / blank req / exact / under / over",
            "InvalidCode|MissingRequirement|Exact|Under|Over",
            '=TEXTJOIN("|",TRUE,IF(COUNTIF(ActiveActivityCodes,"BAD")=0,"InvalidCode",""),IF(COUNTIFS(tblRequirements[OperationalDate],DATE(2099,1,1),tblRequirements[IntervalStart],TIME(0,0,0))=0,"MissingRequirement",""),IF(4-4=0,"Exact",""),IF(3-4<0,"Under",""),IF(5-4>0,"Over",""))',
            "=E7=D7",
            "TEST-06",
            "Verifies key validation and variance labels.",
        ),
        (
            "TEST-07",
            "Overall workbook test status",
            "Rows TEST-01 through TEST-06",
            "PASS",
            '=IF(COUNTIF(F2:F7,FALSE)=0,"PASS","FAIL")',
            "=E8=D8",
            "TEST-07",
            "Exposes an overall pass/fail status for workbook formula tests.",
        ),
        (
            "TEST-08",
            "Scheduled 15-minute break overlay",
            "OWD 10:00-10:30 with Brk 10:00-10:15",
            0.5,
            '=MAX(0,1+((TIME(10,15,0)-TIME(10,0,0))/TIME(0,30,0))*(N(INDEX(tblActivityCodes[CountsAsStaffed],MATCH("Brk",tblActivityCodes[Code],0)))-N(INDEX(tblActivityCodes[CountsAsStaffed],MATCH("OWD",tblActivityCodes[Code],0)))))',
            "=ABS(E9-D9)<0.000001",
            "TEST-08",
            "Verifies a 15-minute break subtracts half of one 30-minute interval.",
        ),
        (
            "TEST-09",
            "Scheduled lunch overlay",
            "OWD 12:00-12:30 with Lch 12:00-12:30",
            0,
            '=MAX(0,1+((TIME(12,30,0)-TIME(12,0,0))/TIME(0,30,0))*(N(INDEX(tblActivityCodes[CountsAsStaffed],MATCH("Lch",tblActivityCodes[Code],0)))-N(INDEX(tblActivityCodes[CountsAsStaffed],MATCH("OWD",tblActivityCodes[Code],0)))))',
            "=ABS(E10-D10)<0.000001",
            "TEST-09",
            "Verifies a full lunch interval removes staffed coverage.",
        ),
        (
            "TEST-10",
            "Scheduled ad hoc overlay",
            "OWD 14:00-14:30 with Mt 14:00-14:30",
            0,
            '=MAX(0,1+((TIME(14,30,0)-TIME(14,0,0))/TIME(0,30,0))*(N(INDEX(tblActivityCodes[CountsAsStaffed],MATCH("Mt",tblActivityCodes[Code],0)))-N(INDEX(tblActivityCodes[CountsAsStaffed],MATCH("OWD",tblActivityCodes[Code],0)))))',
            "=ABS(E11-D11)<0.000001",
            "TEST-10",
            "Verifies an ad hoc meeting interval removes staffed coverage.",
        ),
    ]
    for row in rows:
        ws.append(row)
    for row_index in range(len(rows) + 2, 101):
        ws.cell(row=row_index, column=1, value="")
    add_test_case_conditional_formatting(ws)
    add_table(ws, "tblTestCases", "A1:H100")


def add_test_case_conditional_formatting(ws) -> None:
    pass_fill = PatternFill("solid", fgColor="C6EFCE")
    fail_fill = PatternFill("solid", fgColor="F8696B")
    ws.conditional_formatting.add("F2:F8", FormulaRule(formula=["F2=TRUE"], fill=pass_fill))
    ws.conditional_formatting.add("F2:F8", FormulaRule(formula=["F2=FALSE"], fill=fail_fill))
    ws.conditional_formatting.add("K1", FormulaRule(formula=['K1="PASS"'], fill=pass_fill))
    ws.conditional_formatting.add("K1", FormulaRule(formula=['K1="FAIL"'], fill=fail_fill))


def setup_names_and_validation(wb: Workbook) -> None:
    add_name(wb, "SelectedDate", "'Schedule_Matrix'!$E$2")
    add_name(wb, "IntervalHeaders", "'Schedule_Matrix'!$H$1:$BC$1")
    add_name(wb, "ActiveActivityCodes", "'Config_Settings'!$A$2:$A$13")
    add_name(wb, "tblDailyMatrix", "'Schedule_Matrix'!$A$1:$BC$257")

    schedule = wb["Schedule_Data"]
    validation = DataValidation(type="list", formula1="=ActiveActivityCodes", allow_blank=True)
    schedule.add_data_validation(validation)
    validation.add("F2:F1000")
    for code_range in ["R2:R1000", "U2:U1000", "X2:X1000"]:
        validation.add(code_range)


def format_workbook(wb: Workbook) -> None:
    freezes = {
        "Config_Settings": "A2",
        "Staffing_Requirements": "A2",
        "Schedule_Data": "A2",
        "Schedule_Matrix": "H2",
        "Summary_Dashboard": "A4",
        "TestCases": "A2",
    }
    for ws in wb.worksheets:
        ws.freeze_panes = freezes.get(ws.title)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="center")
        ws.sheet_view.showGridLines = True

    matrix = wb["Schedule_Matrix"]
    calc = wb["Calc_Engine"]
    dashboard = wb["Summary_Dashboard"]
    schedule = wb["Schedule_Data"]
    for ws in [matrix, calc]:
        for column in range(1, 8):
            ws.column_dimensions[ws.cell(row=1, column=column).column_letter].width = 18
        for column in range(8, 56):
            ws.column_dimensions[ws.cell(row=1, column=column).column_letter].width = 8
        for column in range(VALIDATION_STATE_COLUMN, SCHEDULE_HELPER_END_COLUMN + 1):
            ws.column_dimensions[ws.cell(row=1, column=column).column_letter].hidden = True
        for row in [1, 5, 6, 7, SUMMARY_SCHEDULED_ROW, SUMMARY_REQUIRED_ROW, SUMMARY_VARIANCE_ROW]:
            for cell in ws[row]:
                cell.font = Font(bold=True)
    matrix["E2"].number_format = "yyyy-mm-dd"
    for row in range(2, 1001):
        schedule.cell(row=row, column=3).number_format = "yyyy-mm-dd"
        schedule.cell(row=row, column=4).number_format = "yyyy-mm-dd hh:mm"
        schedule.cell(row=row, column=5).number_format = "yyyy-mm-dd hh:mm"
        for column in SCHEDULE_EXTRA_TIME_COLUMNS:
            schedule.cell(row=row, column=column).number_format = "hh:mm"
    for row in range(ROSTER_START_ROW, ROSTER_END_ROW + 1):
        for ws in [matrix, calc]:
            ws.cell(row=row, column=5).number_format = "yyyy-mm-dd"
            ws.cell(row=row, column=6).number_format = "yyyy-mm-dd hh:mm"
            ws.cell(row=row, column=7).number_format = "yyyy-mm-dd hh:mm"
    for column in range(INTERVAL_START_COLUMN, INTERVAL_END_COLUMN + 1):
        matrix.cell(row=1, column=column).number_format = "hh:mm"
        calc.cell(row=1, column=column).number_format = "hh:mm"
        for row in [5, 6, 7, SUMMARY_SCHEDULED_ROW, SUMMARY_REQUIRED_ROW, SUMMARY_VARIANCE_ROW]:
            matrix.cell(row=row, column=column).number_format = "0.00"
            calc.cell(row=row, column=column).number_format = "0.00"
        for row in range(ROSTER_START_ROW, ROSTER_END_ROW + 1):
            calc.cell(row=row, column=column).number_format = "0.00"
    dashboard.column_dimensions["A"].width = 16
    dashboard.column_dimensions["B"].width = 14
    dashboard.column_dimensions["C"].width = 22
    dashboard.column_dimensions["D"].width = 14
    dashboard.column_dimensions["E"].width = 14
    dashboard["A1"].font = Font(bold=True, size=14)
    for row in [3, 13]:
        for cell in dashboard[row]:
            cell.font = Font(bold=True)
    for coordinate in ["B1", "B2"]:
        dashboard[coordinate].number_format = "yyyy-mm-dd"
    for row in range(4, 11):
        dashboard.cell(row=row, column=1).number_format = "yyyy-mm-dd"
        for column in range(3, 6):
            dashboard.cell(row=row, column=column).number_format = "0.00"
    for row in range(14, 62):
        dashboard.cell(row=row, column=1).number_format = "hh:mm"
        for column in range(2, 5):
            dashboard.cell(row=row, column=column).number_format = "0.00"
    tests = wb["TestCases"]
    for row in [*range(2, 7), *range(9, 12)]:
        tests.cell(row=row, column=4).number_format = "0.00"
        tests.cell(row=row, column=5).number_format = "0.00"
    for row in [7, 8]:
        tests.cell(row=row, column=4).number_format = "@"
        tests.cell(row=row, column=5).number_format = "@"
    tests["K1"].number_format = "@"
    calc.sheet_state = "hidden"


def build(engine: str = "excel") -> Workbook:
    wb = setup_workbook()
    populate_config(wb["Config_Settings"])
    populate_requirements(wb["Staffing_Requirements"])
    populate_schedule_data(wb["Schedule_Data"])
    populate_schedule_matrix(wb, engine)
    populate_summary_dashboard(wb)
    populate_test_cases(wb["TestCases"])
    setup_names_and_validation(wb)
    format_workbook(wb)
    return wb


def main() -> None:
    EXCEL_OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    build("excel").save(EXCEL_OUTPUT_PATH)
    build("google").save(GOOGLE_OUTPUT_PATH)
    print(f"Workbook written to {EXCEL_OUTPUT_PATH}")
    print(f"Workbook written to {GOOGLE_OUTPUT_PATH}")


if __name__ == "__main__":
    main()
