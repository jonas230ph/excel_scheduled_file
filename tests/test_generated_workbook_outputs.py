"""Rigorous output-file checks for the Excel and Google workbook variants."""

from __future__ import annotations

import unittest
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer

from scripts.build_wfm_workbook import EXCEL_OUTPUT_PATH, GOOGLE_OUTPUT_PATH, SHEETS


ERROR_TOKENS = ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A")


class GeneratedWorkbookOutputTests(unittest.TestCase):
    def assert_workbook_is_openable_and_formula_safe(self, path: Path, engine: str) -> None:
        self.assertTrue(path.exists(), f"Missing workbook: {path}")

        with ZipFile(path) as archive:
            self.assertIsNone(archive.testzip(), f"Corrupt zip package: {path}")

        wb = load_workbook(path, data_only=False)
        self.assertEqual(wb.sheetnames, SHEETS)
        self.assertEqual(wb["Calc_Engine"].sheet_state, "hidden")

        calc_formula = wb["Calc_Engine"]["A8"].value
        self.assertIsInstance(calc_formula, str)
        if engine == "excel":
            self.assertTrue(calc_formula.startswith("=IFERROR(INDEX(Schedule_Data!$A$2:$A$1000,AGGREGATE("))
            self.assertNotIn("FILTER(", calc_formula)
        else:
            self.assertTrue(calc_formula.startswith("=IFERROR(INDEX(FILTER(Schedule_Data!$A$2:$A$1000,"))
            self.assertIn("FILTER(Schedule_Data!$A$2:$A$1000", calc_formula)
            self.assertNotIn("AGGREGATE(", calc_formula)

        formulas = 0
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    value = cell.value
                    if not isinstance(value, str) or not value.startswith("="):
                        continue
                    formulas += 1
                    Tokenizer(value)
                    upper = value.upper()
                    for token in ERROR_TOKENS:
                        self.assertNotIn(token, upper, f"Unexpected error token in {ws.title}!{cell.coordinate}: {value}")

        self.assertGreater(formulas, 1000, f"Expected many formulas in {path}")

        dashboard = wb["Summary_Dashboard"]
        self.assertEqual(dashboard["C4"].value, '=IF(A4=Schedule_Matrix!$E$2,ROUND(SUM(Schedule_Matrix!$H$260:$BC$260),2),0)')
        self.assertEqual(len(dashboard._charts), 2)

        matrix = wb["Schedule_Matrix"]
        self.assertTrue(matrix["H8"].value.startswith("=LET("))
        self.assertIn("$E$2", matrix["H8"].value)

    def test_excel_workbook_opens_and_parses_every_formula(self) -> None:
        self.assert_workbook_is_openable_and_formula_safe(EXCEL_OUTPUT_PATH, "excel")

    def test_google_workbook_opens_and_parses_every_formula(self) -> None:
        self.assert_workbook_is_openable_and_formula_safe(GOOGLE_OUTPUT_PATH, "google")


if __name__ == "__main__":
    unittest.main()
