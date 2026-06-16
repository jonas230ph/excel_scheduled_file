"""Verify the generated Phase 1 WFM workbook structure."""

from __future__ import annotations

from pathlib import Path
import sys

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from scripts.build_wfm_workbook import (
    SCHEDULE_EXTRA_FIELDS,
    SCHEDULE_HELPER_END_COLUMN,
    SCHEDULE_HELPER_START_COLUMN,
    VALIDATION_STATE_COLUMN,
)

WORKBOOK_PATH = ROOT / "excel schedule" / "wfm_scheduling_matrix.xlsx"

EXPECTED_SHEETS = [
    "Config_Settings",
    "Staffing_Requirements",
    "Schedule_Data",
    "Calc_Engine",
    "Schedule_Matrix",
    "Summary_Dashboard",
    "TestCases",
]

CONFIG_HEADERS = ["Code", "DisplayLabel", "IsActive", "CountsAsStaffed", "IsPaid", "Category", "NumericValue", "ColorGroup"]
PARAMETER_HEADERS = ["Parameter", "Value"]
THRESHOLD_HEADERS = ["State", "MinVariance", "MaxVariance", "FillColor", "FontColor"]
REQUIREMENT_HEADERS = ["OperationalDate", "IntervalStart", "RequiredHeadcount"]
SCHEDULE_HEADERS = [
    "EmployeeID",
    "Name",
    "OperationalDate",
    "ShiftStart",
    "ShiftEnd",
    "ActivityCode",
    "ContractualHours",
    "CurrentShift",
    "Notes",
    "ValidationState",
    "SourceRowID",
    *SCHEDULE_EXTRA_FIELDS,
]
MATRIX_HEADERS = ["EmployeeID", "Name", "ContractualHours", "CurrentShift", "TargetDaySelect", "ShiftStart", "ShiftEnd"]
TEST_HEADERS = ["TestID", "Scenario", "Input", "Expected", "Actual", "Pass", "Requirement", "Notes"]
WEEKLY_DASHBOARD_HEADERS = ["Date", "Day", "Scheduled Productive", "Required", "Over/Under"]
INTRADAY_DASHBOARD_HEADERS = ["Interval", "Scheduled Productive", "Required", "Over/Under"]


def row_values(ws, row: int, start_col: int, end_col: int) -> list:
    return [ws.cell(row=row, column=column).value for column in range(start_col, end_col + 1)]


def assert_tables(wb) -> None:
    expected_tables = {
        "Config_Settings": {"tblActivityCodes", "tblParameters", "tblThresholds"},
        "Staffing_Requirements": {"tblRequirements"},
        "Schedule_Data": {"tblScheduleData"},
        "TestCases": {"tblTestCases"},
    }
    for sheet_name, expected in expected_tables.items():
        actual = set(wb[sheet_name].tables.keys())
        missing = expected - actual
        assert not missing, f"{sheet_name} missing tables: {sorted(missing)}"


def assert_defined_names(wb) -> None:
    actual = set(wb.defined_names.keys())
    expected = {"SelectedDate", "ActiveActivityCodes", "IntervalHeaders", "tblDailyMatrix"}
    missing = expected - actual
    assert not missing, f"Missing workbook defined names: {sorted(missing)}"
    assert wb.defined_names["SelectedDate"].attr_text == "'Schedule_Matrix'!$E$2"
    assert wb.defined_names["IntervalHeaders"].attr_text == "'Schedule_Matrix'!$H$1:$BC$1"
    assert wb.defined_names["tblDailyMatrix"].attr_text == "'Schedule_Matrix'!$A$1:$BC$257"


def assert_config(ws) -> None:
    assert row_values(ws, 1, 1, 8) == CONFIG_HEADERS
    assert row_values(ws, 1, 10, 11) == PARAMETER_HEADERS
    assert row_values(ws, 1, 13, 17) == THRESHOLD_HEADERS
    code_rows = {ws.cell(row=row, column=1).value: ws.cell(row=row, column=4).value for row in range(2, 41)}
    for code in ["OWD", "OT", "Brk", "Lch", "Mt", "Trn", "PTO", "Sick", "Off", "UA"]:
        assert code in code_rows, f"Missing activity code {code}"
    assert code_rows["OWD"] is True
    assert code_rows["OT"] is True
    for code in ["Brk", "Lch", "Mt", "Trn"]:
        assert code_rows[code] is False
    threshold_states = {ws.cell(row=row, column=13).value for row in range(2, 21)}
    for state in [
        "Under",
        "Exact",
        "Over",
        "InvalidCode",
        "InvalidSchedule",
        "Duplicate",
        "MissingDate",
        "MissingRequirement",
        "Nonproductive",
    ]:
        assert state in threshold_states, f"Missing threshold state {state}"


def assert_requirements(ws) -> None:
    assert row_values(ws, 1, 1, 3) == REQUIREMENT_HEADERS
    requirement_rows = [row for row in range(2, 338) if ws.cell(row=row, column=1).value]
    assert len(requirement_rows) == 336, f"Expected 336 weekly requirement rows, found {len(requirement_rows)}"
    requirement_dates = {ws.cell(row=row, column=1).value for row in requirement_rows}
    assert len(requirement_dates) == 7, f"Expected 7 requirement dates, found {len(requirement_dates)}"


def assert_schedule_data(ws) -> None:
    assert row_values(ws, 1, 1, 26) == SCHEDULE_HEADERS
    shift_ends = [ws.cell(row=row, column=5).value for row in range(2, 5)]
    assert any(
        hasattr(value, "date") and value.date().isoformat() == "2026-06-12"
        for value in shift_ends
    ), "Missing overnight sample row"
    for coordinate in ["D2", "E2", "D3", "E3", "D4", "E4"]:
        assert ws[coordinate].data_type == "d", f"{coordinate} must be a typed Excel datetime"
        assert ws[coordinate].number_format == "yyyy-mm-dd hh:mm", f"{coordinate} must use datetime format"
    assert ws["J2"].value.startswith('=IF(COUNTA($A2:$F2)=0,""')
    assert '$C2="","MissingDate"' in ws["J2"].value
    assert '$D2=$E2),"InvalidSchedule"' in ws["J2"].value
    assert 'COUNTIF(ActiveActivityCodes,$F2)=0,"InvalidCode"' in ws["J2"].value
    assert ',"Duplicate","Valid"' in ws["J2"].value
    assert ws["L1"].value == "Break1Start"
    assert ws["M1"].value == "Break1End"
    assert ws["N1"].value == "Break2Start"
    assert ws["O1"].value == "Break2End"
    assert ws["P1"].value == "LunchStart"
    assert ws["Q1"].value == "LunchEnd"
    assert ws["R1"].value == "Adhoc1Code"
    assert ws["R2"].value == "Mt"
    assert ws["L2"].number_format == "hh:mm"
    assert ws["P2"].number_format == "hh:mm"
    assert ws["S2"].number_format == "hh:mm"
    validations = list(ws.data_validations.dataValidation)
    assert validations, "Missing activity-code data validation"
    assert any("F2:F1000" in str(validation.sqref) for validation in validations), "Activity validation not applied to F2:F1000"
    for code_range in ["R2:R1000", "U2:U1000", "X2:X1000"]:
        assert any(code_range in str(validation.sqref) for validation in validations), f"Ad hoc validation not applied to {code_range}"
    cf_ranges = {str(item.sqref) for item in ws.conditional_formatting}
    assert "A2:Z1000" in cf_ranges, f"Missing schedule-data validation formatting: {cf_ranges}"


def assert_schedule_matrix(ws) -> None:
    assert row_values(ws, 1, 1, 7) == MATRIX_HEADERS
    interval_headers = row_values(ws, 1, 8, 55)
    assert len(interval_headers) == 48
    assert all(value for value in interval_headers), "Interval headers must be nonblank"
    assert ws["H1"].value == "=TIME(0,0,0)+(COLUMN()-COLUMN($H$1))*TIME(0,30,0)"
    assert ws["BC1"].value == "=TIME(0,0,0)+(COLUMN()-COLUMN($H$1))*TIME(0,30,0)"
    assert ws["G5"].value == "Scheduled Productive"
    assert ws["G6"].value == "Required"
    assert ws["G7"].value == "Over/Under"
    assert ws["BD1"].value == "ValidationState"
    assert ws.column_dimensions["BD"].hidden, "ValidationState helper column must be hidden"
    assert row_values(ws, 1, SCHEDULE_HELPER_START_COLUMN, SCHEDULE_HELPER_END_COLUMN) == SCHEDULE_EXTRA_FIELDS
    for column in range(VALIDATION_STATE_COLUMN, SCHEDULE_HELPER_END_COLUMN + 1):
        letter = ws.cell(row=1, column=column).column_letter
        assert ws.column_dimensions[letter].hidden, f"{letter} helper column must be hidden"
    assert ws["G260"].value == "Scheduled Productive"
    assert ws["G261"].value == "Required"
    assert ws["G262"].value == "Over/Under"
    assert ws["G202"].value == "=Calc_Engine!G202"
    assert ws["H5"].value == "=SUM(Calc_Engine!H$8:H$257)"
    assert ws["H5"].number_format == "0.00"
    assert "COUNTIFS(tblRequirements[OperationalDate],$E$2" in ws["H6"].value
    assert ')=0,"",SUMIFS(tblRequirements[RequiredHeadcount]' in ws["H6"].value
    assert ws["H7"].value == '=IF($E$2="","MissingDate",IF(H$6="","MissingRequirement",ROUND(H$5-H$6,2)))'
    assert ws["H260"].value == "=SUM(Calc_Engine!H$8:H$257)"
    assert ws["H260"].number_format == "0.00"
    assert "COUNTIFS(tblRequirements[OperationalDate],$E$2" in ws["H261"].value
    assert ')=0,"",SUMIFS(tblRequirements[RequiredHeadcount]' in ws["H261"].value
    assert ws["H262"].value == '=IF($E$2="","MissingDate",IF(H$261="","MissingRequirement",ROUND(H$260-H$261,2)))'
    assert ws["H262"].number_format == "0.00"
    assert ws["A8"].value == "=Calc_Engine!A8"
    assert ws["F8"].value == "=Calc_Engine!F8"
    assert ws["G8"].value == "=Calc_Engine!G8"
    assert ws["BD8"].value == "=Calc_Engine!BD8"
    assert ws["BE8"].value == "=Calc_Engine!BE8"
    assert ws["BS8"].value == "=Calc_Engine!BS8"
    assert "Break1Start" in ws["H8"].value
    assert "Adhoc3Code" in ws["H8"].value
    cf_ranges = {str(item.sqref) for item in ws.conditional_formatting}
    assert "A8:BC257" in cf_ranges, f"Missing selected-row validation formatting: {cf_ranges}"
    assert "E2" in cf_ranges, f"Missing selected-date validation formatting: {cf_ranges}"
    assert "H8:BC257" in cf_ranges, f"Missing 250-row body conditional formatting: {cf_ranges}"
    assert "H262:BC262" in cf_ranges, f"Missing moved variance conditional formatting: {cf_ranges}"
    assert len(ws.conditional_formatting) > 0, "Missing conditional formatting"


def assert_calc_engine(ws) -> None:
    assert ws["A8"].value.startswith("=IFERROR(INDEX(FILTER(")
    assert "SelectedDate" not in ws["A8"].value
    assert "LET(" not in ws["A8"].value
    assert "tblScheduleData[OperationalDate]=SelectedDate" not in ws["A8"].value
    assert "(INT(tblScheduleData[ShiftStart])=0)*tblScheduleData[OperationalDate]" in ws["A8"].value
    assert "(INT(tblScheduleData[ShiftEnd])=0)*tblScheduleData[OperationalDate]" in ws["A8"].value
    assert 'tblScheduleData[EmployeeID]<>""' in ws["A8"].value
    assert 'Schedule_Matrix!$E$2<>""' in ws["A8"].value
    assert "<Schedule_Matrix!$E$2+1" in ws["A8"].value
    assert ">Schedule_Matrix!$E$2" in ws["A8"].value
    assert ws["D8"].value.startswith("=IFERROR(INDEX(FILTER(")
    assert ws["F8"].value.startswith("=IFERROR(INDEX(FILTER(")
    assert ws["BD1"].value == "ValidationState"
    assert ws["BD8"].value.startswith("=IFERROR(INDEX(FILTER(")
    assert "FILTER(tblScheduleData[ValidationState]" in ws["BD8"].value
    assert ws.column_dimensions["BD"].hidden, "Calc validation helper column must be hidden"
    assert row_values(ws, 1, SCHEDULE_HELPER_START_COLUMN, SCHEDULE_HELPER_END_COLUMN) == SCHEDULE_EXTRA_FIELDS
    for column in range(VALIDATION_STATE_COLUMN, SCHEDULE_HELPER_END_COLUMN + 1):
        letter = ws.cell(row=1, column=column).column_letter
        assert ws.column_dimensions[letter].hidden, f"Calc {letter} helper column must be hidden"
    assert "FILTER(tblScheduleData[Break1Start]" in ws["BE8"].value
    assert "FILTER(tblScheduleData[Adhoc3End]" in ws["BS8"].value
    assert ws["H8"].value.startswith("=LET(")
    assert "SelectedDate" not in ws["H8"].value
    assert "Schedule_Matrix!$E$2" in ws["H8"].value
    assert "overlayCoverage" in ws["H8"].value
    assert ws["H8"].number_format == "0.00"


def assert_test_cases(ws) -> None:
    assert row_values(ws, 1, 1, 8) == TEST_HEADERS
    populated = [row for row in range(2, 101) if ws.cell(row=row, column=1).value]
    assert len(populated) >= 10, f"Expected at least 10 test case rows, found {len(populated)}"
    expected_ids = [f"TEST-{index:02d}" for index in range(1, 11)]
    assert [ws.cell(row=row, column=1).value for row in range(2, 12)] == expected_ids
    assert ws["J1"].value == "OverallStatus"
    assert ws["K1"].value == '=IF(COUNTIF(F2:F11,FALSE)=0,"PASS","FAIL")'
    for row in range(2, 12):
        actual = ws.cell(row=row, column=5).value
        passed = ws.cell(row=row, column=6).value
        requirement = ws.cell(row=row, column=7).value
        assert isinstance(actual, str) and actual.startswith("="), f"E{row} must contain a workbook-native formula"
        assert isinstance(passed, str) and passed.startswith("="), f"F{row} must contain a pass/fail formula"
        assert requirement == f"TEST-{row - 1:02d}", f"G{row} requirement trace is wrong"
    for row in [*range(2, 7), *range(9, 12)]:
        assert ws.cell(row=row, column=5).number_format == "0.00", f"E{row} must render numeric test output as a number"
    assert "09:00-17:00" in ws["C2"].value
    assert "22:00-06:00" in ws["C3"].value
    assert "08:15" in ws["C4"].value
    assert 'MATCH("OWD",tblActivityCodes[Code],0)' in ws["E5"].value
    assert all(code in ws["E6"].value for code in ["Brk", "Lch", "Mt", "Trn"])
    for state in ["InvalidCode", "MissingRequirement", "Exact", "Under", "Over"]:
        assert state in ws["E7"].value, f"TEST-06 missing {state}"
    assert ws["E8"].value == '=IF(COUNTIF(F2:F7,FALSE)=0,"PASS","FAIL")'
    assert "Brk" in ws["E9"].value and "OWD" in ws["E9"].value
    assert "Lch" in ws["E10"].value and "OWD" in ws["E10"].value
    assert "Mt" in ws["E11"].value and "OWD" in ws["E11"].value


def assert_summary_dashboard(ws) -> None:
    assert ws["A1"].value == "Weekly Staffing Summary"
    assert ws["B1"].value == "=Schedule_Matrix!$E$2-WEEKDAY(Schedule_Matrix!$E$2,2)+1"
    assert ws["A2"].value == "Selected Date"
    assert ws["B2"].value == "=Schedule_Matrix!$E$2"
    assert row_values(ws, 3, 1, 5) == WEEKLY_DASHBOARD_HEADERS
    assert ws["A4"].value == "=$B$1"
    assert ws["A5"].value == "=A4+1"
    assert ws["B4"].value == '=TEXT(A4,"ddd")'
    assert ws["C4"].value == '=IF(A4=Schedule_Matrix!$E$2,ROUND(SUM(Schedule_Matrix!$H$260:$BC$260),2),0)'
    assert "Schedule_Matrix!$H$260:$BC$260" in ws["C4"].value
    assert "SUMIF(tblActivityCodes[Code],baseCode" not in ws["C4"].value
    assert "SUMPRODUCT(" not in ws["C4"].value
    assert "MMULT(" not in ws["C4"].value
    assert ws["D4"].value == "=SUMIFS(tblRequirements[RequiredHeadcount],tblRequirements[OperationalDate],A4)"
    assert ws["E4"].value == '=IF(D4="","",ROUND(C4-D4,2))'
    assert row_values(ws, 13, 1, 4) == INTRADAY_DASHBOARD_HEADERS
    assert ws["A14"].value == "=Schedule_Matrix!H$1"
    assert ws["B14"].value == "=Schedule_Matrix!H$260"
    assert ws["C14"].value == "=Schedule_Matrix!H$261"
    assert ws["D14"].value == "=Schedule_Matrix!H$262"
    assert len(ws._charts) == 2, "Summary_Dashboard must contain weekly and intraday charts"


def main() -> None:
    assert WORKBOOK_PATH.exists(), f"Workbook missing: {WORKBOOK_PATH}"
    assert WORKBOOK_PATH.suffix == ".xlsx", "Workbook must be a macro-free .xlsx file"
    wb = load_workbook(WORKBOOK_PATH, data_only=False)
    assert wb.sheetnames == EXPECTED_SHEETS, f"Unexpected sheet order: {wb.sheetnames}"
    assert wb["Calc_Engine"].sheet_state == "hidden"
    assert_config(wb["Config_Settings"])
    assert_requirements(wb["Staffing_Requirements"])
    assert_schedule_data(wb["Schedule_Data"])
    assert_calc_engine(wb["Calc_Engine"])
    assert_schedule_matrix(wb["Schedule_Matrix"])
    assert_summary_dashboard(wb["Summary_Dashboard"])
    assert_test_cases(wb["TestCases"])
    assert_tables(wb)
    assert_defined_names(wb)
    print("WFM workbook verification passed")


if __name__ == "__main__":
    main()
